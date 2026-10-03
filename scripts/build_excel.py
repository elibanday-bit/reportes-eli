"""Genera el modelo Excel con fórmulas (Resumen pool, Detalle asesor, Parametros) sobre la base.
Uso: python scripts/build_excel.py
Después, abrir el archivo en Excel para que calcule (openpyxl no guarda valores calculados)."""
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1]
BASE=str(ROOT/'data'/'ANALISIS_ESTADO-PROYECCION.xlsx')
OUT=str(ROOT/'output'/'Estatus_Proyeccion_Comisiones_S37-S40.xlsx')
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.formatting.rule import CellIsRule, DataBarRule
from openpyxl.comments import Comment

wb=openpyxl.load_workbook(BASE)
base=wb['base']
N=264  # last data row
F=Font(name='Arial',size=10); FB=Font(name='Arial',size=10,bold=True)
FH=Font(name='Arial',size=10,bold=True,color='FFFFFF'); FBL=Font(name='Arial',size=10,color='0000FF')
FG=Font(name='Arial',size=10,color='008000')
HF=PatternFill('solid',fgColor='1F3A5F'); YF=PatternFill('solid',fgColor='FFFF00'); GF=PatternFill('solid',fgColor='EEF2F6')
thin=Side(style='thin',color='C9D1DA'); BD=Border(top=thin,bottom=thin,left=thin,right=thin)
CEN=Alignment(horizontal='center',vertical='center',wrap_text=True)

# ---------- Parametros ----------
p=wb.create_sheet('Parametros',0)
p['A1']='Parámetros del modelo de proyección'; p['A1'].font=Font(name='Arial',size=13,bold=True)
rows=[('Peso PM (UG+WA+CAD+RN)',0.7,'Dato del usuario (columna PESO PM de la base)'),
      ('Peso CM (MOD B)',0.3,'Dato del usuario (columna PESO CM de la base)'),
      ('Umbral "cerca de activar"',0.7,'Supuesto: asesores entre 70% y 80% se muestran como recuperables')]
for i,(a,b,c) in enumerate(rows,start=3):
    p.cell(i,1,a).font=F; x=p.cell(i,2,b); x.font=FBL; x.fill=YF; x.number_format='0%'; p.cell(i,3,c).font=Font(name='Arial',size=9,italic=True,color='555555')
notes=['Supuestos de la proyección:',
 '1. El periodo medido es S37–S40 (las columnas de vacaciones y el HC de la base cubren esas 4 semanas).',
 '2. UG, WA, CAD y RN se consideran cerrados a S40 (con dato cargado en las 4 semanas).',
 '3. MOD B (CM) solo tiene carga completa en S37–S38; S39 es parcial y S40 no tiene dato.',
 '   Se proyecta con el ritmo de S37–S38 por semana activa × disponibilidad (Vac S39, Vac S40). Si el dato real es mayor, se usa el real.',
 '4. Cumplimiento = Peso PM × %PM + Peso CM × %CM. Nivel y acelerador según hoja "tablas" (menos de 80% = No activa).',
 '5. Gasto = CARTA (target prorrateado por % cuota) × acelerador.',
 '6. La columna TOTAL VENTA de la base no cuadra con el detalle semanal; se recalcula desde el detalle.',
 '7. Lima Sur no tiene TARGET ni % cuota cargado: su gasto sale en 0. ReiNew no tiene ventas MOD B: su %CM sale en 0.']
for i,t in enumerate(notes,start=8):
    p.cell(i,1,t).font=FB if i==8 else F
p.column_dimensions['A'].width=34; p.column_dimensions['B'].width=10; p.column_dimensions['C'].width=70
PPM="Parametros!$B$3"; PCM="Parametros!$B$4"; PNEAR="Parametros!$B$5"
TH="tablas!$C$3:$C$8"; LV="tablas!$B$3:$B$8"; AC="tablas!$D$3:$D$8"

# ---------- Detalle ----------
d=wb.create_sheet('Detalle asesor',1)
hdr=['Código PS','Pool','Pool equivalente','Líder grupo','Jefe inmediato','HC',
 'Venta PM','Meta PM','% PM',
 'CM real cargado','Ritmo CM / sem activa','CM S39 proy','CM S40 proy','CM proyectado','Meta CM','% CM real','% CM proy',
 'Cumpl. actual','Cumpl. proyectado','Nivel proyectado','Acelerador proy','Carta (S/)','Gasto proyectado (S/)',
 'Nivel actual','Gasto actual (S/)','Estado','Falta a 80% (pts)']
for j,h in enumerate(hdr,1):
    c=d.cell(1,j,h); c.font=FH; c.fill=HF; c.alignment=CEN; c.border=BD
d.row_dimensions[1].height=42
for r in range(2,N+1):
    b=f"base!"
    fx={
    1:f"={b}A{r}",2:f"={b}E{r}",3:f"={b}D{r}",4:f"={b}F{r}",5:f"={b}G{r}",6:f"=N({b}O{r})",
    7:f"=SUM({b}U{r}:{b}AF{r})+N({b}AK{r})",8:f"=N({b}S{r})",9:f"=IFERROR(G{r}/H{r},0)",
    10:f"=SUM({b}AG{r}:{b}AJ{r})",
    11:f"=IFERROR(({b}AG{r}+{b}AH{r})/({b}K{r}+{b}L{r}),0)",
    12:f"=MAX(N({b}AI{r}),K{r}*N({b}M{r}))",13:f"=MAX(N({b}AJ{r}),K{r}*N({b}N{r}))",
    14:f"={b}AG{r}+{b}AH{r}+L{r}+M{r}",15:f"=N({b}T{r})",
    16:f"=IFERROR(J{r}/O{r},0)",17:f"=IFERROR(N{r}/O{r},0)",
    18:f"={PPM}*I{r}+{PCM}*P{r}",19:f"={PPM}*I{r}+{PCM}*Q{r}",
    20:f'=IF(S{r}<tablas!$C$3,"No activa",INDEX({LV},MATCH(S{r},{TH},1)))',
    21:f"=IF(S{r}<tablas!$C$3,0,INDEX({AC},MATCH(S{r},{TH},1)))",
    22:f"=N({b}AQ{r})",23:f"=U{r}*V{r}",
    24:f'=IF(R{r}<tablas!$C$3,"No activa",INDEX({LV},MATCH(R{r},{TH},1)))',
    25:f"=IF(R{r}<tablas!$C$3,0,INDEX({AC},MATCH(R{r},{TH},1)))*V{r}",
    26:f'=IF(S{r}>=tablas!$C$3,"Activa",IF(S{r}>={PNEAR},"Cerca","No activa"))',
    27:f"=MAX(0,tablas!$C$3-S{r})"}
    for j,v in fx.items():
        c=d.cell(r,j,v); c.font=FG if j<=6 or j in(8,15,22) else F; c.border=BD
        if j in(9,16,17,18,19,27): c.number_format='0.0%'
        elif j in(11,12,13,14): c.number_format='0.0'
        elif j in(22,23,25): c.number_format='#,##0;(#,##0);-'
        elif j==21: c.number_format='0.0x'
widths=[12,17,13,28,28,6,8,8,8,9,9,9,9,10,8,9,9,10,11,11,10,10,12,10,11,10,10]
for j,w in enumerate(widths,1): d.column_dimensions[L(j)].width=w
d.freeze_panes='C2'; d.auto_filter.ref=f"A1:{L(len(hdr))}{N}"
d.conditional_formatting.add(f"Z2:Z{N}",CellIsRule(operator='equal',formula=['"Activa"'],fill=PatternFill('solid',fgColor='D5F0DC')))
d.conditional_formatting.add(f"Z2:Z{N}",CellIsRule(operator='equal',formula=['"Cerca"'],fill=PatternFill('solid',fgColor='FFF1C7')))
d.conditional_formatting.add(f"Z2:Z{N}",CellIsRule(operator='equal',formula=['"No activa"'],fill=PatternFill('solid',fgColor='F8D9D6')))
d['K1'].comment=Comment('(MOD B S37 + S38) / (Vac S37 + Vac S38). Vac = factor de disponibilidad de la semana (1 = semana completa).','Claude')
d['S1'].comment=Comment('Peso PM × %PM + Peso CM × %CM proyectado (ver Parametros).','Claude')

# ---------- Resumen pool ----------
s=wb.create_sheet('Resumen pool',0)
s['A1']='Estatus y proyección por pool — Periodo S37–S40'; s['A1'].font=Font(name='Arial',size=13,bold=True)
s['A2']='Fuente: hoja base (ventas UG/WA/CAD/RN a S40; MOD B cargado a S39 parcial). Supuestos en hoja Parametros.'; s['A2'].font=Font(name='Arial',size=9,italic=True,color='555555')
H=['Pool','Dotación','HC efectivo','Venta PM','Meta PM','% PM','CM real','CM proyectado','Meta CM','% CM proy','Cumpl. actual','Cumpl. proyectado',
   'Desfase PM (und)','Desfase CM (und)','Activan','Cerca (70-80%)','No activan','% activación','Carta (S/)','Gasto actual (S/)','Gasto proyectado (S/)','% gasto / carta']
for j,h in enumerate(H,1):
    c=s.cell(4,j,h); c.font=FH; c.fill=HF; c.alignment=CEN; c.border=BD
s.row_dimensions[4].height=42
pools=['Inbound Premium','Inbound Regular','Outbound Premium','Outbound Regular','Chattigo','Counter','Ate','Lima Sur','ReiNew']
D="'Detalle asesor'!"; rng=lambda c:f"{D}${c}$2:${c}${N}"
r0=5
for i,pl in enumerate(pools):
    r=r0+i
    k=f"$A{r}"
    f={1:pl,2:f"=COUNTIF({rng('B')},{k})",3:f"=SUMIF({rng('B')},{k},{rng('F')})",4:f"=SUMIF({rng('B')},{k},{rng('G')})",
       5:f"=SUMIF({rng('B')},{k},{rng('H')})",6:f"=IFERROR(D{r}/E{r},0)",7:f"=SUMIF({rng('B')},{k},{rng('J')})",
       8:f"=SUMIF({rng('B')},{k},{rng('N')})",9:f"=SUMIF({rng('B')},{k},{rng('O')})",10:f"=IFERROR(H{r}/I{r},0)",
       11:f"={PPM}*F{r}+{PCM}*IFERROR(G{r}/I{r},0)",12:f"={PPM}*F{r}+{PCM}*J{r}",
       13:f"=D{r}-E{r}",14:f"=H{r}-I{r}",
       15:f'=COUNTIFS({rng("B")},{k},{rng("Z")},"Activa")',16:f'=COUNTIFS({rng("B")},{k},{rng("Z")},"Cerca")',
       17:f'=COUNTIFS({rng("B")},{k},{rng("Z")},"No activa")',18:f"=IFERROR(O{r}/B{r},0)",
       19:f"=SUMIF({rng('B')},{k},{rng('V')})",20:f"=SUMIF({rng('B')},{k},{rng('Y')})",21:f"=SUMIF({rng('B')},{k},{rng('W')})",22:f"=IFERROR(U{r}/S{r},0)"}
    for j,v in f.items():
        c=s.cell(r,j,v); c.font=F; c.border=BD
rt=r0+len(pools)
s.cell(rt,1,'TOTAL')
for j in range(2,23):
    col=L(j)
    if j in(6,): v=f"=IFERROR(D{rt}/E{rt},0)"
    elif j==10: v=f"=IFERROR(H{rt}/I{rt},0)"
    elif j==11: v=f"={PPM}*F{rt}+{PCM}*IFERROR(G{rt}/I{rt},0)"
    elif j==12: v=f"={PPM}*F{rt}+{PCM}*J{rt}"
    elif j==18: v=f"=IFERROR(O{rt}/B{rt},0)"
    elif j==22: v=f"=IFERROR(U{rt}/S{rt},0)"
    else: v=f"=SUM({col}{r0}:{col}{rt-1})"
    s.cell(rt,j,v)
for j in range(1,23):
    c=s.cell(rt,j); c.font=FB; c.fill=GF; c.border=BD
for r in range(r0,rt+1):
    for j in (6,10,11,12,18,22): s.cell(r,j).number_format='0.0%'
    for j in (3,8): s.cell(r,j).number_format='0.0'
    for j in (13,14): s.cell(r,j).number_format='#,##0;(#,##0);-'
    for j in (19,20,21): s.cell(r,j).number_format='#,##0;(#,##0);-'
s.conditional_formatting.add(f"L{r0}:L{rt-1}",DataBarRule(start_type='num',start_value=0,end_type='num',end_value=1.3,color='2E7D5B'))
s.column_dimensions['A'].width=18
for j in range(2,23): s.column_dimensions[L(j)].width=11
s.freeze_panes='B5'
s.cell(rt+2,1,'Check: la dotación total debe igualar las filas de la base (263).').font=Font(name='Arial',size=9,italic=True)
s.cell(rt+2,6,f"=IF(B{rt}=COUNTA(base!$A$2:$A${N}),\"OK\",\"Revisar\")").font=FB
wb.move_sheet('Parametros',-(wb.index(wb['Parametros'])-2))
wb.save(OUT)
print(wb.sheetnames)
